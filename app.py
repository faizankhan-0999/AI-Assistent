"""AI Resume ATS Checker
Upload a resume (PDF / DOCX / TXT) and get an ATS score plus improvement tips,
powered by Google Gemini Flash and Streamlit.
"""

import json
import os
import re
from io import BytesIO

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader

# ----------------------------- Configuration -----------------------------
MODEL_NAME = os.getenv("GEMINI_MODEL", "gemini-2.5-flash")
MAX_RESUME_CHARS = 15000      # keep the prompt a reasonable size
MAX_JD_CHARS = 8000
MIN_RESUME_CHARS = 100        # below this, the file is probably scanned/empty

st.set_page_config(page_title="AI Resume ATS Checker", page_icon="📄", layout="wide")


# ----------------------------- Helpers -----------------------------------
def get_api_key() -> str | None:
    """Read the key from Streamlit secrets, then the environment, then the sidebar."""
    try:
        key = st.secrets.get("GEMINI_API_KEY")
        if key:
            return key
    except Exception:
        pass  # no secrets.toml present
    return os.getenv("GEMINI_API_KEY") or st.session_state.get("sidebar_key")


def extract_text(uploaded_file) -> str:
    """Extract plain text from an uploaded PDF, DOCX or TXT file."""
    name = uploaded_file.name.lower()
    data = uploaded_file.getvalue()

    if name.endswith(".pdf"):
        reader = PdfReader(BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ValueError("This PDF is password protected.")
        return "\n".join((page.extract_text() or "") for page in reader.pages).strip()

    if name.endswith(".docx"):
        doc = Document(BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:  # many resumes use tables for layout
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        return "\n".join(p for p in parts if p.strip()).strip()

    if name.endswith(".txt"):
        return data.decode("utf-8", errors="ignore").strip()

    raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")


def build_prompt(resume_text: str, job_description: str) -> str:
    jd_block = (
        f"JOB DESCRIPTION:\n{job_description[:MAX_JD_CHARS]}\n\n"
        "Score keyword match against this job description."
        if job_description.strip()
        else "No job description was given. Judge the resume against general "
        "ATS best practices for the role the resume appears to target."
    )
    return f"""You are an expert ATS (Applicant Tracking System) analyst and resume coach.
Analyze the resume below and respond ONLY with valid JSON in exactly this shape:

{{
  "ats_score": <integer 0-100>,
  "summary": "<2-3 sentence overall assessment>",
  "section_scores": {{
    "formatting": <0-100>,
    "keywords": <0-100>,
    "experience_impact": <0-100>,
    "skills": <0-100>,
    "education": <0-100>,
    "readability": <0-100>
  }},
  "strengths": ["<short point>", "..."],
  "weaknesses": ["<short point>", "..."],
  "missing_keywords": ["<keyword>", "..."],
  "improvements": [
    {{"priority": "High|Medium|Low", "section": "<resume section>",
      "issue": "<what is wrong>", "fix": "<specific action to take>"}}
  ],
  "rewrite_examples": [
    {{"before": "<weak bullet from the resume>", "after": "<stronger rewritten bullet>"}}
  ]
}}

Scoring guidance: be honest and calibrated. Most resumes score 50-80. Penalize
missing contact info, no quantified achievements, weak action verbs, missing
sections, keyword gaps, tables/columns/graphics that break ATS parsing, and
overly long or dense text. Give 5-8 improvements and 2-3 rewrite examples that
are based on real lines from the resume. Do not invent experience.

{jd_block}

RESUME:
{resume_text[:MAX_RESUME_CHARS]}
"""


def parse_json_response(text: str) -> dict:
    """Parse model output into a dict, tolerating markdown fences or extra text."""
    text = (text or "").strip()
    text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text, flags=re.IGNORECASE)
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", text, flags=re.DOTALL)
        if match:
            return json.loads(match.group(0))
        raise ValueError("The AI did not return valid JSON. Please try again.")


def clamp_score(value) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError):
        return 0


def analyze_resume(resume_text: str, job_description: str, api_key: str) -> dict:
    client = genai.Client(api_key=api_key)
    response = client.models.generate_content(
        model=MODEL_NAME,
        contents=build_prompt(resume_text, job_description),
        config=types.GenerateContentConfig(
            temperature=0.2,
            response_mime_type="application/json",
        ),
    )
    return parse_json_response(response.text)


def score_label(score: int) -> tuple[str, str]:
    if score >= 80:
        return "Excellent", "🟢"
    if score >= 65:
        return "Good", "🟡"
    if score >= 50:
        return "Needs work", "🟠"
    return "Poor", "🔴"


# ----------------------------- UI ----------------------------------------
def render_results(result: dict) -> None:
    score = clamp_score(result.get("ats_score"))
    label, icon = score_label(score)

    col1, col2 = st.columns([1, 3])
    with col1:
        st.metric("ATS Score", f"{score}/100")
        st.markdown(f"### {icon} {label}")
    with col2:
        st.progress(score / 100)
        st.write(result.get("summary", ""))

    sections = result.get("section_scores") or {}
    if sections:
        st.subheader("Score breakdown")
        cols = st.columns(len(sections))
        for col, (name, val) in zip(cols, sections.items()):
            col.metric(name.replace("_", " ").title(), clamp_score(val))

    left, right = st.columns(2)
    with left:
        st.subheader("✅ Strengths")
        for item in result.get("strengths", []):
            st.markdown(f"- {item}")
    with right:
        st.subheader("⚠️ Weaknesses")
        for item in result.get("weaknesses", []):
            st.markdown(f"- {item}")

    missing = result.get("missing_keywords", [])
    if missing:
        st.subheader("🔑 Missing keywords")
        st.write(" ".join(f"`{k}`" for k in missing))

    improvements = result.get("improvements", [])
    if improvements:
        st.subheader("🛠️ Recommended improvements")
        order = {"high": 0, "medium": 1, "low": 2}
        improvements = sorted(
            improvements, key=lambda x: order.get(str(x.get("priority", "")).lower(), 3)
        )
        for imp in improvements:
            prio = str(imp.get("priority", "")).title()
            dot = {"High": "🔴", "Medium": "🟠", "Low": "🟢"}.get(prio, "⚪")
            with st.expander(f"{dot} {prio} — {imp.get('section', 'General')}"):
                st.markdown(f"**Issue:** {imp.get('issue', '')}")
                st.markdown(f"**Fix:** {imp.get('fix', '')}")

    examples = result.get("rewrite_examples", [])
    if examples:
        st.subheader("✍️ Rewrite examples")
        for ex in examples:
            st.markdown(f"**Before:** {ex.get('before', '')}")
            st.markdown(f"**After:** {ex.get('after', '')}")
            st.divider()

    st.download_button(
        "Download report (JSON)",
        data=json.dumps(result, indent=2),
        file_name="ats_report.json",
        mime="application/json",
    )


def main() -> None:
    st.title("📄 AI Resume ATS Checker")
    st.caption(f"Upload your resume to get an ATS score and tips to improve it. Model: {MODEL_NAME}")

    with st.sidebar:
        st.header("Settings")
        has_key = False
        try:
            has_key = bool(st.secrets.get("GEMINI_API_KEY"))
        except Exception:
            pass
        if not (has_key or os.getenv("GEMINI_API_KEY")):
            st.text_input(
                "Gemini API key",
                type="password",
                key="sidebar_key",
                help="Get a free key at https://aistudio.google.com/apikey",
            )
        else:
            st.success("API key loaded from secrets.")
        st.info("Your resume is sent to Google's Gemini API for analysis and is not stored by this app.")

    uploaded = st.file_uploader("Upload your resume", type=["pdf", "docx", "txt"])
    job_description = st.text_area(
        "Job description (optional, improves keyword matching)",
        height=150,
        placeholder="Paste the job posting here...",
    )

    if st.button("Analyze resume", type="primary", disabled=uploaded is None):
        api_key = get_api_key()
        if not api_key:
            st.error("Please provide a Gemini API key in the sidebar or in secrets.")
            return

        try:
            with st.spinner("Reading your resume..."):
                text = extract_text(uploaded)
            if len(text) < MIN_RESUME_CHARS:
                st.error(
                    "Could not extract enough text. The file may be a scanned image. "
                    "Please upload a text-based PDF or a DOCX file."
                )
                return
            with st.spinner("Analyzing with Gemini..."):
                result = analyze_resume(text, job_description, api_key)
            st.session_state["result"] = result
        except Exception as exc:  # show a friendly message instead of a traceback
            st.error(f"Something went wrong: {exc}")
            return

    if "result" in st.session_state:
        render_results(st.session_state["result"])


if __name__ == "__main__":
    main()
